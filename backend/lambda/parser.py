import io


def parse_resume(body: bytes, content_type: str, filename: str = "") -> tuple[str, str | None]:
    """
    アップロードされたファイルからテキストと写真（base64 data URL）を抽出する。
    写真抽出はdocxのみ対応。戻り値: (text, photo_b64 or None)
    """
    ct = content_type.lower()
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if "pdf" in ct or ext == "pdf":
        return _parse_pdf(body), None
    elif "wordprocessingml" in ct or "docx" in ct or ext == "docx":
        return _parse_docx(body)
    elif "spreadsheetml" in ct or ext == "xlsx":
        return _parse_xlsx(body), None
    else:
        return _parse_txt(body), None


def _parse_pdf(body: bytes) -> str:
    import pdfplumber

    pages = []
    with pdfplumber.open(io.BytesIO(body)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                pages.append(text.strip())
    return "\n\n".join(pages)


def _parse_docx(body: bytes) -> tuple[str, str | None]:
    import base64
    from docx import Document

    doc = Document(io.BytesIO(body))
    lines = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)

    # テーブル内テキストも抽出
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    # ドキュメント内の最初の画像をbase64 data URLで抽出
    photo_b64 = None
    try:
        for rel in doc.part.rels.values():
            if "image" in rel.reltype:
                image_part = rel.target_part
                mime = image_part.content_type  # e.g. "image/jpeg"
                img_bytes = image_part.blob
                b64 = base64.b64encode(img_bytes).decode("ascii")
                photo_b64 = f"data:{mime};base64,{b64}"
                break
    except Exception as e:
        print(f"[WARN] docx image extraction failed: {e}")

    return "\n".join(lines), photo_b64


def _parse_xlsx(body: bytes) -> str:
    from openpyxl import load_workbook

    # read_only=True でメモリ効率化、data_only=True で数式ではなく値を取得
    wb = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    sections = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            sections.append(f"【{sheet_name}】\n" + "\n".join(rows))

    wb.close()
    return "\n\n".join(sections)


def _parse_txt(body: bytes) -> str:
    # UTF-8 → Shift_JIS の順でデコードを試みる
    for encoding in ("utf-8", "shift_jis", "cp932"):
        try:
            return body.decode(encoding)
        except UnicodeDecodeError:
            continue
    return body.decode("utf-8", errors="replace")
